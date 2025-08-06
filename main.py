import argparse
import asyncio
import glob
import os
import time
from datetime import datetime

import pandas as pd
import pytz
import telegram
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from fetch_data.fetch_data import get_sp500_top200_tickers, get_financial_data
from google_sheet.google_sheet import upload_and_convert_to_gsheet
from signaler.signaler import evaluate_stock

# Telegram bot configuration
TELEGRAM_TOKEN = os.getenv('US_VALUE_INVESTMENT_TELEGRAM_BOT_TOKEN')
CHAT_ID = os.getenv('US_VALUE_INVESTMENT_TELEGRAM_CHAT_ID')
GDRIVE_FOLDER_ID = os.getenv('GDRIVE_FOLDER_ID')
HEARTBEAT_INTERVAL = 7200  # 2 hours in seconds


def signal_description(signal):
  return {
    'BUY': '저평가 - 매수 고려',
    'SELL': '고평가 - 매도 고려',
    'HOLD': '적정가 - 보유 유지'
  }.get(signal, '의견 없음')


def save_colored_excel(df, filename='signals.xlsx'):
  wb = Workbook()
  ws = wb.active

  # 헤더
  ws.append(df.columns.tolist())

  # 색상 정의
  red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC",
                         fill_type="solid")  # BUY
  blue_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF",
                          fill_type="solid")  # SELL

  # Check if 'decision' column exists
  if 'decision' not in df.columns:
    print(
      f"Warning: 'decision' column not found in DataFrame. Saving without coloring.")
    for row in df.itertuples(index=False):
      ws.append(row)
    wb.save(filename)
    print(f"🎨 파일 저장 완료: {filename}")
    return

  # 데이터 행 추가 및 색상 적용
  for i, row in df.iterrows():
    row_data = row.tolist()
    ws.append(row_data)

    decision = row.get('decision', '')
    if decision in ['BUY', 'SELL']:
      try:
        target_cell = ws.cell(row=i + 2,
                              column=df.columns.get_loc('decision') + 1)
        if decision == 'BUY':
          target_cell.fill = red_fill
        elif decision == 'SELL':
          target_cell.fill = blue_fill
      except ValueError as e:
        print(f"Error applying color for row {i + 2}: {e}")

  wb.save(filename)
  print(f"🎨 색상 포함된 파일 저장 완료: {filename}")


def is_us_market_open():
  """Check if current time in KST is within U.S. market hours (9:30 AM to 4:00 PM EDT)."""
  kst = pytz.timezone("Asia/Seoul")
  now = datetime.now(kst)
  current_time = now.time()

  market_open = datetime.strptime("22:30", "%H:%M").time()
  market_close = datetime.strptime("05:00", "%H:%M").time()

  # Market hours span midnight in KST
  if current_time >= market_open or current_time < market_close:
    return True
  return False


async def send_telegram_notification(bot, ticker, signal):
  """Send a notification to Telegram for a given stock signal."""
  message = f"Stock Signal: {ticker} - {signal_description(signal)}"
  await bot.send_message(chat_id=CHAT_ID, text=message)


async def send_heartbeat_notification(bot):
  """Send a heartbeat notification to Telegram with the current KST time."""
  kst = pytz.timezone("Asia/Seoul")
  current_time = datetime.now(kst).strftime('%Y-%m-%d %H:%M:%S KST')
  message = f"Heartbeat: Script is running at {current_time}"
  await bot.send_message(chat_id=CHAT_ID, text=message)
  print(f"Sent heartbeat at {current_time}")


async def send_startup_notification(bot):
  """Send a notification to Telegram when monitoring mode starts."""
  kst = pytz.timezone("Asia/Seoul")
  current_time = datetime.now(kst).strftime('%Y-%m-%d %H:%M:%S KST')
  message = f"Monitoring started at {current_time}"
  await bot.send_message(chat_id=CHAT_ID, text=message)
  print(f"Sent startup notification at {current_time}")


async def monitor_and_notify(output_dir):
  """Monitor the latest signals file for buy/sell signals and send Telegram notifications."""
  bot = telegram.Bot(token=TELEGRAM_TOKEN)
  # Send startup notification
  await send_startup_notification(bot)

  last_modified = 0
  last_heartbeat = time.time()
  base_dir = os.path.dirname(os.path.abspath(__file__))
  output_dir = os.path.join(base_dir, output_dir)

  while True:
    try:
      # Find the latest signals file
      files = glob.glob(os.path.join(output_dir, 'today_signals_*.xlsx'))
      if not files:
        print("No signals file found. Waiting...")
        time.sleep(60)
        continue

      latest_file = max(files, key=os.path.getmtime)
      current_modified = os.path.getmtime(latest_file)

      # Check for heartbeat
      current_time = time.time()
      if current_time - last_heartbeat >= HEARTBEAT_INTERVAL and is_us_market_open():
        await send_heartbeat_notification(bot)
        last_heartbeat = current_time

      if current_modified > last_modified and is_us_market_open():
        df = pd.read_excel(latest_file)
        if 'decision' not in df.columns or 'ticker' not in df.columns:
          print(
            f"Error: 'decision' or 'ticker' column missing in {latest_file}")
          time.sleep(60)
          continue
        for _, row in df.iterrows():
          if row['decision'] in ['BUY', 'SELL']:
            await send_telegram_notification(bot, row['ticker'],
                                             row['decision'])
        last_modified = current_modified
        print(f"Checked {latest_file} for signals.")
      time.sleep(60)  # Check every minute
    except Exception as e:
      print(f"Error in monitoring: {e}")
      time.sleep(60)


def run_pipeline(output_dir):
  """Run the stock evaluation pipeline and save results."""
  results = []
  tickers = get_sp500_top200_tickers()

  for ticker in tickers:
    data = get_financial_data(ticker)
    if not data:
      continue

    signal = evaluate_stock(data)
    data['decision'] = signal if signal else '없음'
    data['comment'] = signal_description(signal)
    results.append(data)

  df = pd.DataFrame(results)
  df = df[['ticker', 'price', 'eps', 'pe', 'pb', 'roe', 'de_ratio', 'decision',
           'comment']]

  # Create output directory
  base_dir = os.path.dirname(os.path.abspath(__file__))
  output_dir = os.path.join(base_dir, output_dir)
  os.makedirs(output_dir, exist_ok=True)

  # Generate filename with timestamp
  today = datetime.today().strftime('%Y-%m-%d_%H-%M-%S')
  filename = os.path.join(output_dir, f'today_signals_{today}.xlsx')

  save_colored_excel(df, filename)
  print(f"✅ 평가 결과 저장 완료 → {filename}")
  return filename


def main():
  parser = argparse.ArgumentParser(
    description="Stock signals script with Google Sheets and Telegram integration")
  parser.add_argument("--mode", choices=["sheets", "monitor"], required=True,
                      help="Mode: 'sheets' to update Google Sheets, 'monitor' for Telegram notifications")
  args = parser.parse_args()

  output_dir = "output_file"

  if args.mode == "sheets":
    filename = run_pipeline(output_dir)
    upload_and_convert_to_gsheet(filename, GDRIVE_FOLDER_ID)
    print(f"✅ Google Sheets 업데이트 완료")
  elif args.mode == "monitor":
    print("Starting monitoring mode...")
    asyncio.run(monitor_and_notify(output_dir))


if __name__ == "__main__":
  main()
